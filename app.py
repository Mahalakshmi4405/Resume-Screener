import os
import re
import json
import uuid
import sqlite3
import csv
import logging
import requests
import threading
from flask import Flask, request, render_template, redirect, url_for, flash, jsonify, send_from_directory
from pypdf import PdfReader
import docx
from bs4 import BeautifulSoup
from werkzeug.utils import secure_filename
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS
from sklearn.metrics.pairwise import cosine_similarity

# Gemini imports
from google import genai

# ---------- CONFIG ----------
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf', 'docx'}
DB_PATH = "resumes.db"
MAX_GEMINI_CHARS = 1600   # limit how much of Gemini response we keep in DB

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler()]
)
logger = logging.getLogger(__name__)

# Flask init
app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET", "supersecret")
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# Gemini client configuration
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
if GEMINI_API_KEY:
    client = genai.Client(api_key=GEMINI_API_KEY)
    GEMINI_ENABLED = True
    logger.info("Gemini client enabled.")
else:
    client = None
    GEMINI_ENABLED = False
    logger.warning("GEMINI_API_KEY not set — Gemini calls will be skipped.")

# ---------- DB INIT & SCHEMA ----------
def init_db():
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS resumes
                 (id TEXT PRIMARY KEY,
                  filename TEXT,
                  resume_text TEXT UNIQUE,
                  job_desc TEXT,
                  job_link TEXT,
                  similarity REAL,
                  gemini_good TEXT,
                  gemini_bad TEXT,
                  gemini_recommendations TEXT,
                  gemini_rating REAL,
                  gemini_match_score REAL,
                  gemini_comparison TEXT,
                  status TEXT,
                  progress INTEGER)''')
    conn.commit()
    conn.close()

init_db()

# ---------- HELPERS ----------
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_text_from_resume(filepath):
    logger.info(f"Extracting text from resume: {filepath}")
    if filepath.lower().endswith(".pdf"):
        try:
            reader = PdfReader(filepath)
            text = ""
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
            return text.strip() or ""
        except Exception as e:
            logger.exception("PDF extraction failed")
            return f"[Error extracting PDF: {e}]"
    elif filepath.lower().endswith((".docx", ".doc")):
        try:
            doc = docx.Document(filepath)
            return "\n".join([para.text for para in doc.paragraphs]) or ""
        except Exception as e:
            logger.exception("DOCX extraction failed")
            return f"[Error extracting DOCX: {e}]"
    return ""

def scrape_text_from_url(url, max_chars=8000):
    logger.info(f"Scraping job description from {url}")
    try:
        res = requests.get(url, timeout=10, headers={"User-Agent": "Mozilla/5.0"})
        res.raise_for_status()
        soup = BeautifulSoup(res.text, "html.parser")
        for script in soup(["script", "style", "noscript"]):
            script.extract()
        text = " ".join(soup.stripped_strings)
        return text[:max_chars]
    except Exception as e:
        logger.exception("Error scraping URL")
        return f"[Error scraping URL: {e}]"

def tokenize_keywords(text, top_k=40):
    text = (text or "").lower()
    words = re.findall(r"\b[a-zA-Z0-9\+\#\.\-]{2,}\b", text)
    filtered = [w for w in words if w not in ENGLISH_STOP_WORDS and not w.isdigit()]
    freq = {}
    for w in filtered:
        freq[w] = freq.get(w, 0) + 1
    sorted_words = sorted(freq.items(), key=lambda x: x[1], reverse=True)
    return [w for w, _ in sorted_words[:top_k]]

def compute_similarity(text1, text2):
    logger.info("Computing similarity between resume and job description...")
    try:
        t1 = (text1 or "").strip()
        t2 = (text2 or "").strip()
        if t1 == "" and t2 == "":
            return 0.0
        vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), max_features=5000)
        tfidf = vectorizer.fit_transform([t1, t2])
        tfidf_sim = cosine_similarity(tfidf[0:1], tfidf[1:2])[0][0]
        kw1 = set(tokenize_keywords(t1, top_k=60))
        kw2 = set(tokenize_keywords(t2, top_k=60))
        jac = len(kw1 & kw2) / len(kw1 | kw2) if kw1 and kw2 else 0.0
        combined = (0.75 * tfidf_sim) + (0.25 * jac)
        sim = round(combined * 100, 2)
        logger.info(f"Similarity score: {sim}%")
        return sim
    except Exception:
        logger.exception("Similarity computation failed")
        return 0.0

def update_db_columns(resume_id, **kwargs):
    if not kwargs:
        return
    keys = list(kwargs.keys())
    placeholders = ", ".join([f"{k} = ?" for k in keys])
    vals = [kwargs[k] for k in keys] + [resume_id]
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    sql = f"UPDATE resumes SET {placeholders} WHERE id = ?"
    c.execute(sql, vals)
    conn.commit()
    conn.close()

def exists_resume_text(text):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute("SELECT id FROM resumes WHERE resume_text = ?", (text,))
    row = c.fetchone()
    conn.close()
    return row is not None

def save_initial_metadata(resume_id, filename, resume_text, job_desc, job_link, similarity=0.0):
    logger.info("Inserting initial metadata row (processing)...")
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    try:
        c.execute("""INSERT INTO resumes
                     (id, filename, resume_text, job_desc, job_link, similarity, status, progress)
                     VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                  (resume_id, filename, resume_text, job_desc, job_link, similarity, "processing", 5))
        conn.commit()
    except sqlite3.IntegrityError:
        logger.warning("Duplicate resume detected while inserting initial metadata.")
        conn.close()
        return False
    conn.close()
    return True

def export_to_csv():
    logger.info("Exporting database to CSV...")
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute("SELECT id, filename, resume_text, job_desc, job_link, similarity, gemini_good, gemini_bad, gemini_recommendations, gemini_rating, gemini_match_score, gemini_comparison FROM resumes")
    rows = c.fetchall()
    conn.close()
    with open("resumes.csv", "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Filename", "Resume_Content", "Job_Description", "Job_Link",
                         "Similarity", "Good_Points", "Bad_Points", "Recommendations",
                         "Rating", "Gemini_Match_Score", "Gemini_Comparison"])
        writer.writerows(rows)
    logger.info("CSV export complete.")

# ---------- Gemini helpers ----------
def gemini_generate_from_text(prompt, model="gemini-2.5-flash"):
    if not GEMINI_ENABLED:
        logger.warning("Gemini not enabled — skipping API call.")
        return "[Gemini disabled: API key not configured]"
    try:
        logger.info("Calling Gemini (text prompt)...")
        response = client.models.generate_content(model=model, contents=[prompt])
        text = getattr(response, "text", "") or str(response)
        logger.debug("Gemini response length: %d", len(text))
        return text
    except Exception as e:
        logger.exception("Gemini API error")
        return f"[Gemini API error: {e}]"

def gemini_field_from_text(resume_text, job_desc_text, instruction_text, expect_json=False, truncate_resume=8000, truncate_jd=8000):
    rtxt = (resume_text or "")[:truncate_resume]
    jtxt = (job_desc_text or "")[:truncate_jd]
    prompt = f"""You are an expert resume screener.
Resume (candidate):
---
{rtxt}
---
Job description:
---
{jtxt}
---
Task: {instruction_text}

Return the requested output. If JSON is requested, return valid JSON only.
"""
    raw = gemini_generate_from_text(prompt)
    if expect_json:
        try:
            parsed = json.loads(raw)
            return parsed
        except Exception:
            m = re.search(r'(\{.*\})', raw, re.DOTALL)
            if m:
                try:
                    parsed = json.loads(m.group(1))
                    return parsed
                except Exception:
                    pass
            return {"raw": raw}
    else:
        return raw.strip()

def normalize_gemini_text(text, max_chars=MAX_GEMINI_CHARS):
    """
    - Trim and sanitize whitespace
    - Convert long single-line outputs to bullet-like lines if they look list-y
    - Truncate at word boundary to max_chars
    """
    if text is None:
        return ""
    s = str(text).strip()

    # If raw JSON-like object, pretty-print a short excerpt
    try:
        parsed = json.loads(s)
        # If it's a list, join with bullets
        if isinstance(parsed, list):
            lines = []
            for item in parsed[:20]:
                lines.append(f"• {str(item).strip()}")
            out = "\n".join(lines)
        elif isinstance(parsed, dict):
            lines = []
            for k, v in list(parsed.items())[:20]:
                lines.append(f"• {k}: {str(v).strip()}")
            out = "\n".join(lines)
        else:
            out = str(parsed)
    except Exception:
        out = s

    # Replace multiple blank lines with one
    out = re.sub(r'\r\n|\r', '\n', out)
    out = re.sub(r'\n{2,}', '\n', out)

    # If text is a single long line with separators like ; or . then attempt split into bullets
    if '\n' not in out and len(out) > 180:
        # split on sentence boundaries or semicolons
        pieces = re.split(r'[;\.\n]\s+', out)
        if len(pieces) > 1:
            lines = [p.strip() for p in pieces if p.strip()]
            out = "\n".join([("• " + l) for l in lines[:30]])
    # Ensure bullets have nice bullet char
    out = re.sub(r'^\s*-\s*', '• ', out, flags=re.M)
    out = re.sub(r'^\s*•\s*', '• ', out, flags=re.M)

    # Truncate safely (word boundary)
    if len(out) > max_chars:
        truncated = out[:max_chars]
        # try truncate at last newline or space
        last_nl = truncated.rfind('\n')
        if last_nl > max_chars - 300:
            truncated = truncated[:last_nl]
        else:
            last_space = truncated.rfind(' ')
            if last_space > 0:
                truncated = truncated[:last_space]
        out = truncated.rstrip() + "\n…"

    return out

# ---------- BACKGROUND PROCESS ----------
def process_resume_background(resume_id, filepath, resume_text, job_desc, job_link):
    try:
        update_db_columns(resume_id, status="processing", progress=10)
        # re-extract (safety)
        try:
            extracted = extract_text_from_resume(filepath)
            if extracted and (not resume_text or len(resume_text.strip()) < 10):
                resume_text = extracted
                update_db_columns(resume_id, resume_text=resume_text)
        except Exception:
            pass

        # local similarity
        update_db_columns(resume_id, progress=20)
        try:
            sim = compute_similarity(resume_text, job_desc)
            update_db_columns(resume_id, similarity=sim, progress=30)
        except Exception:
            sim = 0.0
            update_db_columns(resume_id, similarity=sim, progress=30)

        # Gemini strengths
        update_db_columns(resume_id, progress=40)
        try:
            strengths_raw = gemini_field_from_text(resume_text, job_desc, "List the top 3 strengths of the resume for this job. Keep each point short and specific (max 12 words). Format as bullet points (•).", expect_json=False)
            strengths = normalize_gemini_text(strengths_raw)
            update_db_columns(resume_id, gemini_good=strengths, progress=50)
        except Exception:
            update_db_columns(resume_id, gemini_good="[Error generating strengths]", progress=50)

        # Gemini weaknesses
        update_db_columns(resume_id, progress=55)
        try:
            weaknesses_raw = gemini_field_from_text(resume_text, job_desc, "List up to 3 weaknesses or issues in the resume. Keep each point short and specific (max 12 words). Format as bullet points (•).", expect_json=False)
            weaknesses = normalize_gemini_text(weaknesses_raw)
            update_db_columns(resume_id, gemini_bad=weaknesses, progress=65)
        except Exception:
            update_db_columns(resume_id, gemini_bad="[Error generating weaknesses]", progress=65)

        # Gemini recommendations
        update_db_columns(resume_id, progress=70)
        try:
            recs_raw = gemini_field_from_text(resume_text, job_desc, "Suggest 3 very short, actionable improvements for this resume. Use concise bullet points (•), focused on making the resume stronger for the given job.", expect_json=False)
            recs = normalize_gemini_text(recs_raw)
            update_db_columns(resume_id, gemini_recommendations=recs, progress=75)
        except Exception:
            update_db_columns(resume_id, gemini_recommendations="[Error generating recommendations]", progress=75)

        # Gemini rating
        update_db_columns(resume_id, progress=80)
        try:
            rating_text = gemini_field_from_text(resume_text, job_desc, "Rate the resume for this job on a scale from 0–10. Return ONLY the number.", expect_json=False)
            rating_val = 0.0
            try:
                m = re.search(r"(\d+(\.\d+)?)", str(rating_text))
                if m:
                    rating_val = float(m.group(1))
            except Exception:
                rating_val = 0.0
            update_db_columns(resume_id, gemini_rating=rating_val, progress=85)
        except Exception:
            update_db_columns(resume_id, gemini_rating=0.0, progress=85)

        # Match and comparison (expect JSON)
        update_db_columns(resume_id, progress=88)
        try:
            match_json = gemini_field_from_text(resume_text, job_desc, "Evaluate how well the resume matches the job description. Return JSON only:\\n{\\n  \"match_score\": <0-100>,\\n  \"comparison\": \"2–4 sentence summary comparing resume to job requirements\"\\n}", expect_json=True)
            match_score = 0.0
            comparison_text = ""
            if isinstance(match_json, dict):
                match_score = float(match_json.get("match_score") or match_json.get("match") or match_json.get("score") or 0)
                comparison_text = str(match_json.get("comparison") or match_json.get("summary") or match_json.get("raw") or "")
            else:
                # fallback: treat as text
                try:
                    text = str(match_json)
                    m = re.search(r'(\d+(\.\d+)?)\s*%', text)
                    if m: match_score = float(m.group(1))
                    comparison_text = text
                except Exception:
                    comparison_text = str(match_json)

            combined = match_score or compute_similarity(resume_text, job_desc)
            comparison_text = normalize_gemini_text(comparison_text, max_chars=1200)
            update_db_columns(resume_id, gemini_match_score=combined, gemini_comparison=comparison_text, progress=95)
        except Exception:
            update_db_columns(resume_id, gemini_match_score=0.0, gemini_comparison="[Error]", progress=95)

        update_db_columns(resume_id, status="done", progress=100)
        export_to_csv()
        logger.info(f"Processing complete for {resume_id}")
    except Exception:
        logger.exception("Background processing failed")
        update_db_columns(resume_id, status="failed", progress=0)

# ---------- ROUTES ----------
@app.route("/", methods=["GET", "POST"])
def upload_resume():
    if request.method == "POST":
        logger.info("New resume upload request received.")
        job_desc = request.form.get("job_desc", "").strip()
        job_link = request.form.get("job_link", "").strip()
        resume_file = request.files.get("resume_file")
        resume_link = request.form.get("resume_link", "").strip()

        # If job link provided, scrape and use that text
        if job_link:
            scraped = scrape_text_from_url(job_link)
            if scraped.startswith("[Error"):
                flash(f"❌ Failed to scrape job link: {scraped}")
                return redirect(request.url)
            job_desc = scraped

        if not job_desc:
            flash("❌ Please enter a job description (text or job link).")
            return redirect(request.url)

        resume_id = str(uuid.uuid4())
        filename, filepath = None, None

        # Handle resume file upload or link
        if resume_file and resume_file.filename != "" and allowed_file(resume_file.filename):
            safe_name = secure_filename(resume_file.filename)
            filename = f"{resume_id}_{safe_name}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            resume_file.save(filepath)
            logger.info(f"Saved uploaded resume to {filepath}")
        elif resume_link:
            if not (resume_link.lower().endswith(".pdf") or resume_link.lower().endswith(".docx")):
                flash("❌ Only .pdf or .docx resume links are allowed.")
                return redirect(request.url)
            try:
                response = requests.get(resume_link, timeout=15)
                response.raise_for_status()
                ext = resume_link.split(".")[-1].lower()
                filename = f"{resume_id}_link.{ext}"
                filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                with open(filepath, "wb") as f:
                    f.write(response.content)
                logger.info(f"Downloaded resume from link to {filepath}")
            except Exception as e:
                logger.exception("Error fetching resume link")
                flash(f"❌ Error fetching resume link: {e}")
                return redirect(request.url)
        else:
            flash("❌ Please upload a resume file or provide a resume link.")
            return redirect(request.url)

        # Extract resume text quickly for duplicate check
        resume_text = extract_text_from_resume(filepath).strip()
        if resume_text.startswith("[Error"):
            flash(f"❌ {resume_text}")
            return redirect(request.url)

        if exists_resume_text(resume_text):
            flash("⚠️ Duplicate resume detected (text already in DB). Skipping insert.")
            return redirect(request.url)

        # Save initial DB row with small progress (background will do heavy lift)
        ok = save_initial_metadata(resume_id, filename, resume_text, job_desc, job_link or None, similarity=0.0)
        if not ok:
            flash("⚠️ Duplicate detected or insert failed.")
            return redirect(request.url)

        # Start background processing thread (non-blocking)
        t = threading.Thread(target=process_resume_background, args=(resume_id, filepath, resume_text, job_desc, job_link), daemon=True)
        t.start()

        # Return the upload page but include resume_id so the page JS will start polling
        return render_template("upload.html", resume_id=resume_id)

    # GET
    return render_template("upload.html", resume_id="")

@app.route("/status/<resume_id>")
def status(resume_id):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute("SELECT id, filename, similarity, gemini_good, gemini_bad, gemini_recommendations, gemini_rating, gemini_match_score, gemini_comparison, status, progress FROM resumes WHERE id = ?", (resume_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "not found"}), 404

    (rid, filename, similarity, gemini_good, gemini_bad, gemini_recommendations, gemini_rating, gemini_match_score, gemini_comparison, status_val, progress) = row
    # Ensure values are serializable
    return jsonify({
        "id": rid,
        "filename": filename,
        "similarity": similarity,
        "gemini_good": gemini_good,
        "gemini_bad": gemini_bad,
        "gemini_recommendations": gemini_recommendations,
        "gemini_rating": gemini_rating,
        "gemini_match_score": gemini_match_score,
        "gemini_comparison": gemini_comparison,
        "status": status_val,
        "progress": progress
    })

@app.route("/scrape", methods=["POST"])
def scrape():
    data = request.get_json() or {}
    url = data.get("url", "").strip()
    if not url:
        return jsonify({"content": "", "error": "No URL provided"}), 400
    content = scrape_text_from_url(url)
    return jsonify({"content": content})

@app.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=False)

@app.route("/report/<resume_id>")
def report(resume_id):
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    c.execute("SELECT id, filename, resume_text, job_desc, job_link, similarity, gemini_good, gemini_bad, gemini_recommendations, gemini_rating, gemini_match_score, gemini_comparison, status, progress FROM resumes WHERE id = ?", (resume_id,))
    row = c.fetchone()
    conn.close()
    if not row:
        flash("Report not found.")
        return redirect(url_for("upload_resume"))

    (rid, filename, resume_text, job_desc, job_link, similarity,
     gemini_good, gemini_bad, gemini_recommendations, gemini_rating,
     gemini_match_score, gemini_comparison, status, progress) = row

    resume_snippet = (resume_text or "")[:600] + ("..." if len(resume_text or "") > 600 else "")
    job_snippet = (job_desc or "")[:1000] + ("..." if len(job_desc or "") > 1000 else "")

    if not GEMINI_ENABLED:
        if not gemini_good: gemini_good = "[Gemini disabled]"
        if not gemini_bad: gemini_bad = "[Gemini disabled]"
        if not gemini_recommendations: gemini_recommendations = "[Gemini disabled]"
        if not gemini_comparison: gemini_comparison = "[Gemini disabled]"

    return render_template("report.html",
                           id=rid,
                           filename=filename,
                           resume_snippet=resume_snippet,
                           job_snippet=job_snippet,
                           job_link=job_link,
                           similarity=similarity,
                           gemini_good=gemini_good,
                           gemini_bad=gemini_bad,
                           gemini_recommendations=gemini_recommendations,
                           gemini_rating=gemini_rating,
                           gemini_match_score=gemini_match_score,
                           gemini_comparison=gemini_comparison,
                           status=status,
                           progress=progress)

if __name__ == "__main__":
    logger.info("Starting Flask server...")
    app.run(debug=True)
